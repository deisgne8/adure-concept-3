"""Normalize ADURE inventory workbooks and optionally import them into WordPress."""

from __future__ import annotations

import argparse
import base64
import json
import re
import tempfile
import urllib.error
import urllib.request
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HEADER_ALIASES = {
    "unitname": "unit_code",
    "unitcode": "unit_code",
    "unit": "unit_code",
    "buildingname": "building_name",
    "floor": "floor",
    "floornumber": "floor",
    "areasqmt": "area",
    "areasqm": "area",
    "area": "area",
    "commonarea": "common_area",
    "netarea": "net_area",
    "unittype": "unit_type",
    "unitsubtype": "unit_subtype",
    "unitview": "view",
    "view": "view",
    "status": "status",
    "rent": "rent",
    "totalrentrate": "rent",
    "pricepersqm": "price_per_sqm",
    "annualrentm2maxaed": "price_per_sqm",
    "balcony": "balcony",
    "bathroomno": "bathrooms",
    "maidroom": "maid_room",
    "storeroom": "store_room",
    "studyroom": "study_room",
    "facility": "facilities",
    "ewmeterno": "meter_number",
}

BUILDING_ALIASES = {
    "jasmin tower": "Jasmine Tower",
    "jasmine tower": "Jasmine Tower",
    "al hili tower b": "Hili Tower B",
    "hili tower b": "Hili Tower B",
    "48 burj gate dubai": "48 Burj Gate",
    "48 burj gate": "48 Burj Gate",
    "al manhal": "Al Manhal Tower",
    "al manhal tower": "Al Manhal Tower",
    "qaryat al hidd garden 5": "Garden Residence 5",
    "qaryat al hidd sunrise 1": "Sunrise Residence 1",
    "qaryat al hidd sunrise 2": "Sunrise Residence 2",
    "qaryat al hidd sunrise 3": "Sunrise Residence 3",
    "qaryat al hidd sunrise 4": "Sunrise Residence 4",
    "qaryat al hidd sunrise 5": "Sunrise Residence 5",
}

AMENITY_ALIASES = {
    "beach": "Beach",
    "pool": "Swimming Pool",
    "swimming pool": "Swimming Pool",
    "gym": "Gym",
    "kids playing area": "Kids' Play Area",
    "yoga room": "Yoga Room",
    "facility room": "Facility Room",
    "underground parking": "Underground Parking",
    "built in cooker and washing machine": "Built-in Cooker and Washing Machine",
    "driver room": "Driver's Room",
    "driver's room": "Driver's Room",
    "private garden": "Private Garden",
    "covered parking": "Covered Parking",
    "jacuzzi": "Jacuzzi",
}


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def slugify(value: str) -> str:
    value = value.lower().strip()
    value = re.sub(r"[^a-z0-9]+", "-", value)
    return value.strip("-")


def normalize_header(value: Any) -> str:
    compact = re.sub(r"[^a-z0-9]", "", text(value).lower())
    return HEADER_ALIASES.get(compact, compact)


def canonical_building(value: str) -> str:
    cleaned = re.sub(r"\s+", " ", value).strip(" ,-_")
    return BUILDING_ALIASES.get(cleaned.lower(), cleaned)


def number(value: Any) -> float | None:
    if value is None or text(value).lower() in {"", "n/a", "na", "none", "-"}:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d[\d,]*(?:\.\d+)?", text(value))
    return float(match.group(0).replace(",", "")) if match else None


def whole_number(value: Any) -> int | None:
    parsed = number(value)
    return int(parsed) if parsed is not None else None


def boolean(value: Any) -> bool:
    return text(value).lower() in {"1", "yes", "y", "true", "available"}


def bedrooms(unit_type: str, subtype: str) -> int | None:
    value = f"{unit_type} {subtype}".lower()
    if "studio" in value:
        return 0
    match = re.search(r"(\d+)\s*(?:bed|br)", value)
    return int(match.group(1)) if match else None


def availability(value: Any) -> str:
    value = text(value).lower()
    if value in {"empty", "vacant", "available"}:
        return "available"
    if value in {"rented", "occupied"}:
        return "occupied"
    return "unknown"


def unit_amenities(record: dict[str, Any]) -> list[str]:
    values: list[str] = []
    raw_facilities = text(record.get("facilities"))
    if raw_facilities.lower() not in {"", "no", "n/a", "na", "none", "-"}:
        for raw_value in re.split(r"\s+-\s+|,", raw_facilities):
            normalized = AMENITY_ALIASES.get(raw_value.strip().lower())
            if normalized and normalized not in values:
                values.append(normalized)

    structured = {
        "balcony": "Balcony",
        "maid_room": "Maid Room",
        "store_room": "Store Room",
        "study_room": "Study Room",
    }
    for field, label in structured.items():
        if boolean(record.get(field)) and label not in values:
            values.append(label)
    return values


def source_location(path: Path) -> list[str]:
    name = path.stem.lower()
    if "qaryat al hidd" in name:
        return ["Abu Dhabi", "Qaryat Al Hidd"]
    for location in ("Abu Dhabi", "Al Ain", "Dubai"):
        if location.lower() in name:
            return [location]
    return []


def workbook_paths(source: Path, extraction: Path) -> list[Path]:
    if source.suffix.lower() == ".zip":
        with zipfile.ZipFile(source) as archive:
            archive.extractall(extraction)
        return sorted(extraction.rglob("*.xlsx"))
    if source.is_dir():
        return sorted(source.rglob("*.xlsx"))
    return [source]


def normalize(source: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    buildings: dict[str, dict[str, Any]] = {}
    units: dict[str, dict[str, Any]] = {}
    rejected: list[dict[str, Any]] = []
    duplicates: list[dict[str, Any]] = []
    source_rows = 0

    with tempfile.TemporaryDirectory(prefix="adure-property-import-") as temp_dir:
        paths = workbook_paths(source, Path(temp_dir))
        for path in paths:
            is_retail = "retail" in " ".join(part.lower() for part in path.parts)
            sector = "Retail" if is_retail else "Residential"
            workbook_location = source_location(path) if not is_retail else []
            workbook = load_workbook(path, read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                try:
                    headers = [normalize_header(value) for value in next(rows)]
                except StopIteration:
                    continue
                for row_number, values in enumerate(rows, start=2):
                    if not any(value not in (None, "") for value in values):
                        continue
                    source_rows += 1
                    record = {
                        headers[index]: value
                        for index, value in enumerate(values)
                        if index < len(headers) and headers[index]
                    }
                    building_name = canonical_building(
                        text(record.get("building_name")) or sheet.title
                    )
                    unit_code = text(record.get("unit_code"))
                    source_ref = f"{path.name} / {sheet.title} / row {row_number}"
                    if not building_name or not unit_code:
                        rejected.append(
                            {
                                "source": source_ref,
                                "reason": "Missing building name or unit code",
                            }
                        )
                        continue

                    building_key = slugify(building_name)
                    building = buildings.setdefault(
                        building_key,
                        {
                            "key": building_key,
                            "slug": building_key,
                            "name": building_name,
                            "sectors": [],
                            "location": workbook_location,
                        },
                    )
                    if sector not in building["sectors"]:
                        building["sectors"].append(sector)
                    if workbook_location and not building["location"]:
                        building["location"] = workbook_location

                    unit_key = f"{building_key}:{unit_code.lower()}"
                    if unit_key in units:
                        duplicates.append(
                            {
                                "source": source_ref,
                                "key": unit_key,
                                "reason": "Duplicate building and unit code",
                            }
                        )
                        continue

                    unit_type = text(record.get("unit_type"))
                    subtype = text(record.get("unit_subtype"))
                    source_area = text(record.get("area"))
                    units[unit_key] = {
                        "key": unit_key,
                        "buildingKey": building_key,
                        "buildingName": building_name,
                        "buildingSlug": building_key,
                        "unitCode": unit_code,
                        "sector": sector,
                        "location": workbook_location,
                        "transaction": "lease",
                        "status": availability(record.get("status")),
                        "floor": text(record.get("floor")),
                        "unitType": unit_type,
                        "subtype": subtype,
                        "bedrooms": bedrooms(unit_type, subtype),
                        "bathrooms": whole_number(record.get("bathrooms")),
                        "areaSqm": number(record.get("area")),
                        "commonAreaSqm": number(record.get("common_area")),
                        "netAreaSqm": number(record.get("net_area")),
                        "sourceArea": source_area,
                        "view": text(record.get("view")),
                        "balcony": boolean(record.get("balcony")),
                        "maidRoom": boolean(record.get("maid_room")),
                        "storeRoom": boolean(record.get("store_room")),
                        "studyRoom": boolean(record.get("study_room")),
                        "facilities": text(record.get("facilities")),
                        "amenities": unit_amenities(record),
                        "annualRent": number(record.get("rent")),
                        "pricePerSqm": number(record.get("price_per_sqm")),
                        "meterNumber": text(record.get("meter_number")),
                        "source": source_ref,
                    }
            workbook.close()

    status_counts: dict[str, int] = defaultdict(int)
    sector_counts: dict[str, int] = defaultdict(int)
    amenity_counts: dict[str, int] = defaultdict(int)
    for unit in units.values():
        status_counts[unit["status"]] += 1
        sector_counts[unit["sector"]] += 1
        for amenity in unit["amenities"]:
            amenity_counts[amenity] += 1

    payload = {"buildings": list(buildings.values()), "units": list(units.values())}
    report = {
        "sourceRows": source_rows,
        "buildings": len(buildings),
        "unitsReady": len(units),
        "rejected": rejected,
        "duplicates": duplicates,
        "reconciledRows": len(units) + len(rejected) + len(duplicates),
        "statusCounts": dict(status_counts),
        "sectorCounts": dict(sector_counts),
        "amenityCounts": dict(sorted(amenity_counts.items())),
    }
    return payload, report


def post_batch(
    endpoint: str,
    username: str,
    password: str,
    payload: dict[str, Any],
) -> dict[str, Any]:
    request = urllib.request.Request(
        endpoint,
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": "Basic "
            + base64.b64encode(f"{username}:{password}".encode()).decode(),
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            return json.load(response)
    except urllib.error.HTTPError as error:
        raise RuntimeError(error.read().decode("utf-8")) from error


def import_payload(
    payload: dict[str, Any],
    wordpress_url: str,
    username: str,
    password: str,
    batch_size: int,
) -> dict[str, Any]:
    endpoint = wordpress_url.rstrip("/") + "/wp-json/adure/v1/property-import"
    responses: list[dict[str, Any]] = []
    buildings = payload["buildings"]
    units = payload["units"]
    for start in range(0, len(buildings), batch_size):
        responses.append(
            post_batch(
                endpoint,
                username,
                password,
                {"buildings": buildings[start : start + batch_size]},
            )
        )
    for start in range(0, len(units), batch_size):
        responses.append(
            post_batch(
                endpoint,
                username,
                password,
                {"units": units[start : start + batch_size]},
            )
        )
    responses.append(
        post_batch(endpoint, username, password, {"finalize": True})
    )
    errors = [error for response in responses for error in response.get("errors", [])]
    return {"batches": len(responses), "errors": errors, "locked": responses[-1].get("locked")}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, default=Path("qa/property-import.json"))
    parser.add_argument("--report", type=Path, default=Path("qa/property-import-report.json"))
    parser.add_argument("--wordpress-url")
    parser.add_argument("--username")
    parser.add_argument("--application-password")
    parser.add_argument("--batch-size", type=int, default=50)
    args = parser.parse_args()

    payload, report = normalize(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    args.report.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))

    credentials = (args.wordpress_url, args.username, args.application_password)
    if any(credentials) and not all(credentials):
        raise SystemExit("WordPress URL, username and application password are all required.")
    if all(credentials):
        result = import_payload(
            payload,
            args.wordpress_url,
            args.username,
            args.application_password,
            args.batch_size,
        )
        print(json.dumps(result, indent=2))
        if result["errors"]:
            raise SystemExit(1)


if __name__ == "__main__":
    main()
